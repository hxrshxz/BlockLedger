#!/usr/bin/env python3
"""Generate the SIH 2026 presentation for Team Mercury / BlockLedger."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------------------------------------------------------------- palette
BG = RGBColor(0x0B, 0x0D, 0x10)
CARD = RGBColor(0x14, 0x18, 0x1D)
LINE = RGBColor(0x2A, 0x31, 0x3A)
FG = RGBColor(0xF2, 0xF4, 0xF6)
MUTED = RGBColor(0x9A, 0xA4, 0xB0)
ACCENT = RGBColor(0x4A, 0xE3, 0xB5)
WARN = RGBColor(0xFF, 0x9F, 0x45)
ARROW = RGBColor(0x55, 0x60, 0x6C)

HASHES = ["0x9f3a1c", "0x4b8e07", "0xc21df5", "0x7a60be", "0x1e94d2"]

TITLE_FONT = "Arial"
BODY_FONT = "Arial"

W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def box(
    s,
    x,
    y,
    w,
    h,
    text,
    size=18,
    color=FG,
    bold=False,
    align=PP_ALIGN.LEFT,
    font=BODY_FONT,
    spacing=1.25,
):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def card(s, x, y, w, h, fill=CARD, line=LINE):
    sh = s.shapes.add_shape(5, x, y, w, h)  # rounded rectangle
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.06
    except Exception:
        pass
    return sh


def rule(s, x, y, w, color=ACCENT, thick=Pt(3)):
    sh = s.shapes.add_shape(1, x, y, w, thick)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def header(s, kicker, title):
    box(
        s,
        Inches(0.9),
        Inches(0.55),
        Inches(11.5),
        Inches(0.3),
        kicker.upper(),
        size=12,
        color=ACCENT,
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        Inches(0.9),
        Inches(0.92),
        Inches(11.5),
        Inches(0.7),
        title,
        size=32,
        color=FG,
        bold=True,
        font=TITLE_FONT,
        spacing=1.0,
    )
    rule(s, Inches(0.9), Inches(1.72), Inches(1.1))


def footer(s, n):
    box(
        s,
        Inches(0.9),
        Inches(6.92),
        Inches(6.0),
        Inches(0.3),
        "BlockLedger  ·  Team Mercury  ·  SIH 2026  ·  PS 26125",
        size=10,
        color=MUTED,
        spacing=1.0,
    )
    box(
        s,
        Inches(11.3),
        Inches(6.92),
        Inches(1.1),
        Inches(0.3),
        str(n),
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        spacing=1.0,
    )


def bullets(s, x, y, w, items, size=16, gap=Inches(0.52), dot=ACCENT):
    for i, it in enumerate(items):
        yy = y + gap * i
        d = s.shapes.add_shape(9, x, yy + Inches(0.07), Inches(0.11), Inches(0.11))
        d.fill.solid()
        d.fill.fore_color.rgb = dot
        d.line.fill.background()
        d.shadow.inherit = False
        box(
            s,
            x + Inches(0.28),
            yy,
            w - Inches(0.28),
            Inches(0.45),
            it,
            size=size,
            color=MUTED,
            spacing=1.15,
        )


# ============================================================ 1. TITLE
s = slide()
rule(s, Inches(0.9), Inches(1.65), Inches(1.6))
box(
    s,
    Inches(0.9),
    Inches(1.05),
    Inches(11),
    Inches(0.4),
    "SMART INDIA HACKATHON 2026",
    size=13,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)
box(
    s,
    Inches(0.9),
    Inches(1.95),
    Inches(11.5),
    Inches(1.1),
    "BlockLedger",
    size=66,
    color=FG,
    bold=True,
    font=TITLE_FONT,
    spacing=1.0,
)
box(
    s,
    Inches(0.9),
    Inches(3.02),
    Inches(10.6),
    Inches(0.9),
    "Blockchain-Based Secure Platform for Identity,\nAccess Control and Digital Asset Management",
    size=22,
    color=MUTED,
    spacing=1.25,
)

meta = [
    ("PS ID", "26125"),
    ("Organisation", "Bharat Electronics Limited"),
    ("Category", "Software"),
    ("Theme", "Blockchain & Cybersecurity"),
]
x = Inches(0.9)
for k, v in meta:
    c = card(s, x, Inches(4.35), Inches(2.75), Inches(1.0))
    box(
        s,
        x + Inches(0.25),
        Inches(4.55),
        Inches(2.3),
        Inches(0.25),
        k.upper(),
        size=10,
        color=MUTED,
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        x + Inches(0.25),
        Inches(4.83),
        Inches(2.35),
        Inches(0.4),
        v,
        size=13,
        color=FG,
        bold=True,
        spacing=1.05,
    )
    x += Inches(2.95)

box(
    s,
    Inches(0.9),
    Inches(5.75),
    Inches(11),
    Inches(0.7),
    "Team Mercury  ·  Netaji Subhas University of Technology",
    size=16,
    color=FG,
    bold=True,
    spacing=1.2,
)
box(
    s,
    Inches(0.9),
    Inches(6.12),
    Inches(11),
    Inches(0.4),
    "Harsh (Lead)  ·  Anushka  ·  Sanyyam  ·  Prakhar  ·  Zoya  ·  Shauryan",
    size=13,
    color=MUTED,
    spacing=1.2,
)

# ============================================================ 2. PROBLEM
s = slide()
header(s, "The problem", "Centralised registers cannot prove their own integrity")
probs = [
    (
        "01",
        "Single point of failure",
        "One compromised directory administrator can mint, escalate or delete any identity. There is no independent record, because the same administrator controls the log.",
    ),
    (
        "02",
        "Identity theft and impersonation",
        "Credentials are bearer secrets held centrally. Whoever obtains the store obtains every identity in it. Employees hold no portable, provable identity.",
    ),
    (
        "03",
        "Unverifiable asset provenance",
        "A drawing, firmware image or certificate can be modified after approval, and a downstream consumer has no cryptographic way to prove which bytes were approved.",
    ),
]
y = Inches(2.15)
for num, t, d in probs:
    card(s, Inches(0.9), y, Inches(11.5), Inches(1.32))
    box(
        s,
        Inches(1.25),
        y + Inches(0.3),
        Inches(0.7),
        Inches(0.5),
        num,
        size=26,
        color=ACCENT,
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        Inches(2.15),
        y + Inches(0.22),
        Inches(9.9),
        Inches(0.35),
        t,
        size=17,
        color=FG,
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        Inches(2.15),
        y + Inches(0.62),
        Inches(9.9),
        Inches(0.6),
        d,
        size=12,
        color=MUTED,
        spacing=1.2,
    )
    y += Inches(1.52)

box(
    s,
    Inches(0.9),
    Inches(6.45),
    Inches(11.5),
    Inches(0.4),
    "Net effect: an auditor cannot independently prove that a log has not been altered after the fact.",
    size=14,
    color=WARN,
    bold=True,
    spacing=1.1,
)
footer(s, 2)

# ============================================================ 3. SOLUTION
s = slide()
header(s, "Proposed solution", "Replace the trusted record with cryptographic evidence")
sols = [
    (
        "Decentralised identity",
        "Every actor is a DID — did:blkl:sol:<pubkey> — with a W3C DID Document controlled by a keypair, not a directory row. Pinned to IPFS, so it is content-addressed.",
    ),
    (
        "NFT-based asset ownership",
        "Every asset is a token (BLKL-000123) whose metadata commits to a SHA-256 digest of the file. Change one byte and the digest no longer matches.",
    ),
    (
        "Contract-enforced access control",
        "Privileged actions pass a permission check that reverts with an explicit reason. A caller cannot bypass a check by ignoring a return value.",
    ),
    (
        "Tamper-evident audit trail",
        "Each entry commits to the hash of its predecessor, so retroactive edits are detectable in a single pass — and anchored to Solana devnet.",
    ),
]
positions = [
    (Inches(0.9), Inches(2.15)),
    (Inches(6.95), Inches(2.15)),
    (Inches(0.9), Inches(4.5)),
    (Inches(6.95), Inches(4.5)),
]
for (px, py), (t, d) in zip(positions, sols):
    card(s, px, py, Inches(5.45), Inches(2.05))
    rule(s, px + Inches(0.35), py + Inches(0.38), Inches(0.55))
    box(
        s,
        px + Inches(0.35),
        py + Inches(0.58),
        Inches(4.8),
        Inches(0.4),
        t,
        size=17,
        color=FG,
        bold=True,
        spacing=1.05,
    )
    box(
        s,
        px + Inches(0.35),
        py + Inches(1.02),
        Inches(4.8),
        Inches(0.9),
        d,
        size=12,
        color=MUTED,
        spacing=1.25,
    )
footer(s, 3)

# ============================================================ 4. ARCHITECTURE
s = slide()
header(s, "Architecture", "How a request flows through the system")
layers = [
    ("User + Solana Wallet", "Phantom / Solflare", ACCENT),
    (
        "Next.js Frontend — App Router, React 19",
        "dashboard · identity · assets · access-control · audit · wallet",
        FG,
    ),
    ("Context Layer", "Auth · Chain · Ledger · Storage", FG),
    (
        "Contract Layer",
        "AccessControl · IdentityRegistry · AssetNFT  —  revert on failure",
        FG,
    ),
    (
        "Blockchain Core",
        "DID derivation · SHA-256 hash chaining · append-only ledger + verifier",
        FG,
    ),
]
y = Inches(2.1)
for i, (t, d, col) in enumerate(layers):
    card(s, Inches(0.9), y, Inches(7.4), Inches(0.78))
    box(
        s,
        Inches(1.25),
        y + Inches(0.13),
        Inches(6.8),
        Inches(0.3),
        t,
        size=14,
        color=col,
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        Inches(1.25),
        y + Inches(0.43),
        Inches(6.9),
        Inches(0.28),
        d,
        size=10,
        color=MUTED,
        spacing=1.0,
    )
    if i < len(layers) - 1:
        ar = s.shapes.add_shape(
            67, Inches(4.45), y + Inches(0.79), Inches(0.3), Inches(0.15)
        )
        ar.fill.solid()
        ar.fill.fore_color.rgb = LINE
        ar.line.fill.background()
        ar.shadow.inherit = False
    y += Inches(0.94)

card(s, Inches(8.7), Inches(2.1), Inches(3.7), Inches(1.85))
box(
    s,
    Inches(9.05),
    Inches(2.35),
    Inches(3.1),
    Inches(0.3),
    "IPFS  ·  PINATA",
    size=12,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)
box(
    s,
    Inches(9.05),
    Inches(2.72),
    Inches(3.15),
    Inches(1.05),
    "DID Documents and NFT metadata pinned to real, content-addressed CIDs.",
    size=12,
    color=MUTED,
    spacing=1.25,
)

card(s, Inches(8.7), Inches(4.2), Inches(3.7), Inches(1.85))
box(
    s,
    Inches(9.05),
    Inches(4.45),
    Inches(3.1),
    Inches(0.3),
    "SOLANA DEVNET",
    size=12,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)
box(
    s,
    Inches(9.05),
    Inches(4.82),
    Inches(3.15),
    Inches(1.05),
    "Audit entries anchored as real signed memo transactions, linked to an explorer.",
    size=12,
    color=MUTED,
    spacing=1.25,
)
footer(s, 4)

# ============================================================ 5. AUDIT CHAIN
s = slide()
header(s, "Core innovation", "A log that proves it has not been edited")
box(
    s,
    Inches(0.9),
    Inches(2.1),
    Inches(11.4),
    Inches(0.4),
    "entry[n].hash  =  SHA-256(  entry[n].payload  ||  entry[n-1].hash  )",
    size=19,
    color=ACCENT,
    bold=True,
    font="Consolas",
    spacing=1.0,
)

x = Inches(0.9)
for i in range(5):
    broken = i == 3
    col = WARN if broken else LINE
    card(s, x, Inches(2.85), Inches(1.85), Inches(1.15), line=col)
    box(
        s,
        x + Inches(0.22),
        Inches(3.02),
        Inches(1.5),
        Inches(0.28),
        f"ENTRY {i}",
        size=10,
        color=(WARN if broken else MUTED),
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        x + Inches(0.22),
        Inches(3.32),
        Inches(1.5),
        Inches(0.5),
        ("TAMPERED" if broken else HASHES[i]),
        size=13,
        color=(WARN if broken else FG),
        bold=True,
        font="Consolas",
        spacing=1.0,
    )
    if i < 4:
        ar = s.shapes.add_shape(
            33, x + Inches(1.87), Inches(3.35), Inches(0.38), Inches(0.14)
        )
        ar.fill.solid()
        ar.fill.fore_color.rgb = WARN if i >= 3 else ARROW
        ar.line.fill.background()
        ar.shadow.inherit = False
    x += Inches(2.25)

box(
    s,
    Inches(0.9),
    Inches(4.25),
    Inches(11.4),
    Inches(0.4),
    "verifyChain()  →  { valid: false, brokenAt: 3 }",
    size=16,
    color=WARN,
    bold=True,
    font="Consolas",
    spacing=1.0,
)

box(
    s,
    Inches(0.9),
    Inches(4.95),
    Inches(11.4),
    Inches(0.35),
    "Four classes of tampering are detected, each reporting the exact broken index:",
    size=14,
    color=FG,
    bold=True,
    spacing=1.0,
)
det = ["Field mutation", "Entry reordering", "Entry deletion", "Actor forgery"]
x = Inches(0.9)
for d in det:
    card(s, x, Inches(5.42), Inches(2.75), Inches(0.62))
    box(
        s,
        x,
        Inches(5.6),
        Inches(2.75),
        Inches(0.3),
        d,
        size=13,
        color=FG,
        bold=True,
        align=PP_ALIGN.CENTER,
        spacing=1.0,
    )
    x += Inches(2.95)

box(
    s,
    Inches(0.9),
    Inches(6.3),
    Inches(11.4),
    Inches(0.4),
    "All four are demonstrated live in the prototype.",
    size=13,
    color=MUTED,
    spacing=1.0,
)
footer(s, 5)

# ============================================================ 6. FEATURES
s = slide()
header(s, "Key features", "What the prototype does today")
feats = [
    (
        "Identity",
        "Register, resolve, revoke and reactivate DIDs with W3C DID Documents pinned to IPFS.",
    ),
    (
        "Access control",
        "Four roles mapped through an explicit permission matrix, enforced in UI and contract layer alike.",
    ),
    (
        "Real reverts",
        "Unauthorised calls fail loudly with Solidity-style errors, surfaced directly to the user.",
    ),
    (
        "Digital assets",
        "Mint DID-bound NFTs with content digests, IPFS metadata and full transfer history.",
    ),
    (
        "Audit trail",
        "Hash-chained, append-only log with one-click verification and export.",
    ),
    (
        "On-chain anchor",
        "Signed Solana devnet memo transactions, verifiable in a public block explorer.",
    ),
]
positions = [
    (Inches(0.9), Inches(2.15)),
    (Inches(4.92), Inches(2.15)),
    (Inches(8.94), Inches(2.15)),
    (Inches(0.9), Inches(4.35)),
    (Inches(4.92), Inches(4.35)),
    (Inches(8.94), Inches(4.35)),
]
for (px, py), (t, d) in zip(positions, feats):
    card(s, px, py, Inches(3.45), Inches(1.9))
    rule(s, px + Inches(0.3), py + Inches(0.35), Inches(0.45))
    box(
        s,
        px + Inches(0.3),
        py + Inches(0.55),
        Inches(2.9),
        Inches(0.35),
        t,
        size=16,
        color=FG,
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        px + Inches(0.3),
        py + Inches(0.95),
        Inches(2.9),
        Inches(0.85),
        d,
        size=11,
        color=MUTED,
        spacing=1.25,
    )
footer(s, 6)

# ============================================================ 7. TECH STACK
s = slide()
header(s, "Technology stack", "Built on proven, open components")
groups = [
    (
        "Frontend",
        [
            "Next.js 15 — App Router",
            "React 19",
            "TypeScript",
            "Tailwind CSS v4",
            "shadcn/ui",
            "Framer Motion",
        ],
    ),
    (
        "Blockchain",
        [
            "Solana — devnet",
            "@solana/web3.js",
            "Wallet Adapter",
            "W3C DID",
            "SHA-256 hash chaining",
            "Web Crypto API",
        ],
    ),
    (
        "Storage & Deploy",
        [
            "IPFS",
            "Pinata pinning",
            "Content-addressed CIDs",
            "Vercel",
            "npm",
            "Zero contract deployment",
        ],
    ),
]
x = Inches(0.9)
for title, items in groups:
    card(s, x, Inches(2.15), Inches(3.68), Inches(4.1))
    box(
        s,
        x + Inches(0.35),
        Inches(2.45),
        Inches(3.0),
        Inches(0.35),
        title.upper(),
        size=13,
        color=ACCENT,
        bold=True,
        spacing=1.0,
    )
    bullets(
        s, x + Inches(0.35), Inches(2.95), Inches(3.1), items, size=13, gap=Inches(0.48)
    )
    x += Inches(3.88)
footer(s, 7)

# ============================================================ 8. REAL VS SIMULATED
s = slide()
header(s, "Scope and honesty", "What is real, and what is simulated")
box(
    s,
    Inches(0.9),
    Inches(2.0),
    Inches(11.4),
    Inches(0.4),
    "We state this explicitly rather than overclaim. Nothing in the demo is fabricated.",
    size=14,
    color=MUTED,
    spacing=1.15,
)

card(s, Inches(0.9), Inches(2.6), Inches(5.6), Inches(3.5))
box(
    s,
    Inches(1.25),
    Inches(2.9),
    Inches(4.9),
    Inches(0.35),
    "REAL",
    size=14,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)
bullets(
    s,
    Inches(1.25),
    Inches(3.4),
    Inches(5.0),
    [
        "Solana devnet wallet connection",
        "Signed memo transactions anchoring audit entries",
        "IPFS pinning via Pinata — genuine CIDs",
        "DID derivation and W3C DID Documents",
        "Hash chaining and chain verification",
    ],
    size=12,
    gap=Inches(0.5),
)

card(s, Inches(6.8), Inches(2.6), Inches(5.6), Inches(3.5))
box(
    s,
    Inches(7.15),
    Inches(2.9),
    Inches(4.9),
    Inches(0.35),
    "SIMULATED",
    size=14,
    color=WARN,
    bold=True,
    spacing=1.0,
)
bullets(
    s,
    Inches(7.15),
    Inches(3.4),
    Inches(5.0),
    [
        "Contract layer written in TypeScript, not deployed",
        "Enforcement logic and error semantics mirror an on-chain program",
        "Demo authentication uses seeded credentials",
        "Chosen so the prototype runs with zero deployment cost",
    ],
    size=12,
    gap=Inches(0.62),
    dot=WARN,
)

box(
    s,
    Inches(0.9),
    Inches(6.3),
    Inches(11.4),
    Inches(0.4),
    'Every record carries anchorMode: "onchain" | "simulated" and is badged truthfully in the UI. '
    "No signature or CID is ever fabricated.",
    size=12,
    color=FG,
    bold=True,
    spacing=1.15,
)
footer(s, 8)

# ============================================================ 9. DEMO FLOW
s = slide()
header(s, "Demonstration", "Five steps a reviewer can reproduce")
steps = [
    (
        "1",
        "ADMIN registers an identity",
        "/identity — a DID is issued and its DID Document is pinned to IPFS.",
    ),
    (
        "2",
        "MANAGER mints a digital asset",
        "/assets — metadata pinned, token bound to an owner DID, digest recorded.",
    ),
    (
        "3",
        "AUDITOR is denied a role grant",
        "/access-control — a real contract revert is raised and surfaced.",
    ),
    (
        "4",
        "Chain is verified, then tampered",
        "/audit — the verifier reports the exact index at which the chain breaks.",
    ),
    (
        "5",
        "Entry is anchored on-chain",
        "/wallet — a signed devnet memo transaction, viewable in a block explorer.",
    ),
]
y = Inches(2.15)
for num, t, d in steps:
    card(s, Inches(0.9), y, Inches(11.5), Inches(0.82))
    circ = s.shapes.add_shape(
        9, Inches(1.2), y + Inches(0.21), Inches(0.4), Inches(0.4)
    )
    circ.fill.solid()
    circ.fill.fore_color.rgb = ACCENT
    circ.line.fill.background()
    circ.shadow.inherit = False
    tf = circ.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = num
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = BG
    r.font.name = BODY_FONT
    box(
        s,
        Inches(1.85),
        y + Inches(0.16),
        Inches(4.0),
        Inches(0.35),
        t,
        size=15,
        color=FG,
        bold=True,
        spacing=1.0,
    )
    box(
        s,
        Inches(5.85),
        y + Inches(0.2),
        Inches(6.3),
        Inches(0.45),
        d,
        size=12,
        color=MUTED,
        spacing=1.1,
    )
    y += Inches(0.95)
footer(s, 9)

# ============================================================ 10. IMPACT
s = slide()
header(s, "Impact and future scope", "From prototype to deployment")

box(
    s,
    Inches(0.9),
    Inches(2.05),
    Inches(5.6),
    Inches(0.35),
    "IMPACT",
    size=13,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)
bullets(
    s,
    Inches(0.9),
    Inches(2.55),
    Inches(5.5),
    [
        "Removes the administrator as a single point of trust",
        "Gives auditors mathematical proof of log integrity",
        "Makes asset provenance verifiable by any party",
        "Portable identity that survives directory migration",
        "Independent verification without vendor cooperation",
    ],
    size=13,
    gap=Inches(0.58),
)

box(
    s,
    Inches(6.8),
    Inches(2.05),
    Inches(5.6),
    Inches(0.35),
    "FUTURE SCOPE",
    size=13,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)
bullets(
    s,
    Inches(6.8),
    Inches(2.55),
    Inches(5.5),
    [
        "Deploy the contract layer as Solana programs via Anchor",
        "Wallet-native, signature-based authentication",
        "Verifiable Credentials for cross-organisation attestation",
        "Hardware wallet and HSM-backed keys for ADMIN accounts",
        "Merkle-root batch anchoring to cut cost at scale",
        "Phased LDAP / Active Directory integration",
    ],
    size=13,
    gap=Inches(0.55),
)
footer(s, 10)

# ============================================================ 11. CLOSING
s = slide()
rule(s, Inches(0.9), Inches(2.4), Inches(1.6))
box(
    s,
    Inches(0.9),
    Inches(2.75),
    Inches(11.5),
    Inches(1.0),
    "BlockLedger",
    size=52,
    color=FG,
    bold=True,
    font=TITLE_FONT,
    spacing=1.0,
)
box(
    s,
    Inches(0.9),
    Inches(3.75),
    Inches(10.5),
    Inches(0.6),
    "Identity you own. Access you can prove. Assets you can trace.",
    size=20,
    color=MUTED,
    spacing=1.2,
)

card(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.35))
box(
    s,
    Inches(1.3),
    Inches(4.98),
    Inches(10.5),
    Inches(0.3),
    "TEAM MERCURY",
    size=12,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)
box(
    s,
    Inches(1.3),
    Inches(5.32),
    Inches(10.6),
    Inches(0.35),
    "Harsh (Team Leader)  ·  Anushka  ·  Sanyyam  ·  Prakhar  ·  Zoya  ·  Shauryan",
    size=15,
    color=FG,
    bold=True,
    spacing=1.1,
)
box(
    s,
    Inches(1.3),
    Inches(5.66),
    Inches(10.6),
    Inches(0.3),
    "Netaji Subhas University of Technology  ·  SIH 2026  ·  PS 26125  ·  Bharat Electronics Limited",
    size=11,
    color=MUTED,
    spacing=1.1,
)

box(
    s,
    Inches(0.9),
    Inches(6.35),
    Inches(11.5),
    Inches(0.35),
    "github.com/hxrshxz/BlockLedger",
    size=13,
    color=ACCENT,
    bold=True,
    spacing=1.0,
)

prs.save("submission/Mercury_SIH2026_Presentation.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
